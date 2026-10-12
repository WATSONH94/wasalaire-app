"""
Générateur de la liste nominative CNSS (déclaration mensuelle) — Bénin.
Source : CNSS `cnss.bj/les-listes-nominatives/` — fichier Excel transmis à
carriere@cnss.bj. Colonnes par salarié A-M + ligne récapitulative en tête.

Produit un .xlsx conforme au gabarit et garantit le contrôle de cohérence
(Σ des bruts individuels = masse salariale de la ligne récapitulative).
"""
from __future__ import annotations
from dataclasses import dataclass, field
from datetime import date

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    _OPENPYXL = True
except ImportError:  # dégradation gracieuse → CSV
    _OPENPYXL = False


@dataclass
class LigneSalarie:
    """Une ligne salarié de la liste nominative (colonnes A-M)."""
    num_assurance: str            # B
    nom: str                      # C
    prenoms: str                  # D
    date_embauche: date | None = None   # E
    date_debauchage: date | None = None # F
    jours_ou_heures: float = 0            # G
    salaire_brut: float = 0               # H (assiette CNSS)
    rappel_salaire: float = 0             # I
    periode_rappel: str = ""              # J
    indemnite_retraite: float = 0         # K
    indemnite_licenciement: float = 0     # L
    autres_primes: float = 0              # M


COLONNES = [
    "N° ordre", "N° assurance", "Nom", "Prénoms",
    "Date embauche", "Date débauchage", "Jours/Heures", "Salaire brut",
    "Rappel salaire", "Période rappel", "Indemnité retraite",
    "Indemnité licenciement", "Autres primes",
]


def _fmt_date(d: date | None) -> str:
    return d.strftime("%d/%m/%Y") if d else ""


def controle_coherence(lignes: list[LigneSalarie]) -> dict:
    """Contrôle CNSS : la masse salariale récap = Σ des bruts individuels."""
    masse = sum(l.salaire_brut for l in lignes)
    return {"nb_travailleurs": len(lignes), "masse_salariale": round(masse)}


def generer_liste_nominative(
    lignes: list[LigneSalarie],
    num_employeur: str,
    mois: int,
    annee: int,
    chemin_xlsx: str,
) -> dict:
    """Génère le fichier .xlsx de la liste nominative CNSS.
    Retourne le récap (contrôle de cohérence) et le chemin produit.
    Si openpyxl est absent, produit un CSV équivalent (même chemin en .csv).
    """
    recap = controle_coherence(lignes)

    if not _OPENPYXL:
        chemin = chemin_xlsx.replace(".xlsx", ".csv")
        _generer_csv(lignes, num_employeur, mois, annee, recap, chemin)
        return {**recap, "fichier": chemin, "format": "csv (openpyxl absent)"}

    wb = Workbook()
    ws = wb.active
    ws.title = f"CNSS {mois:02d}-{annee}"

    gras = Font(bold=True)
    entete_fill = PatternFill("solid", fgColor="1F4E78")
    entete_font = Font(bold=True, color="FFFFFF")
    recap_fill = PatternFill("solid", fgColor="FCE4D6")
    bord = Border(*[Side(style="thin", color="BFBFBF")] * 4)
    droite = Alignment(horizontal="right")

    # Ligne 1 : récapitulatif employeur
    ws.append(["N° employeur", num_employeur, "Mois", mois, "Année", annee,
               "Nb travailleurs", recap["nb_travailleurs"],
               "Masse salariale", recap["masse_salariale"]])
    for c in ws[1]:
        c.font = gras
        c.fill = recap_fill

    # Ligne 2 : en-têtes de colonnes A-M
    ws.append(COLONNES)
    for c in ws[2]:
        c.font = entete_font
        c.fill = entete_fill
        c.alignment = Alignment(horizontal="center", wrap_text=True)

    # Lignes salariés
    for i, l in enumerate(lignes, start=1):
        ws.append([
            i, l.num_assurance, l.nom, l.prenoms,
            _fmt_date(l.date_embauche), _fmt_date(l.date_debauchage),
            l.jours_ou_heures, round(l.salaire_brut),
            round(l.rappel_salaire) or "", l.periode_rappel,
            round(l.indemnite_retraite) or "", round(l.indemnite_licenciement) or "",
            round(l.autres_primes) or "",
        ])

    # Ligne total (contrôle)
    ligne_total = ws.max_row + 1
    ws.cell(ligne_total, 7, "TOTAL").font = gras
    cell_total = ws.cell(ligne_total, 8, recap["masse_salariale"])
    cell_total.font = gras
    cell_total.fill = recap_fill

    # Mise en forme
    for col, width in zip("ABCDEFGHIJKLM",
                          [8, 16, 18, 18, 14, 14, 12, 14, 14, 14, 16, 18, 14]):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=3, max_row=ligne_total):
        for c in row:
            c.border = bord
            if isinstance(c.value, (int, float)):
                c.alignment = droite
    ws.freeze_panes = "A3"

    wb.save(chemin_xlsx)
    return {**recap, "fichier": chemin_xlsx, "format": "xlsx",
            "destinataire": "carriere@cnss.bj",
            "source": "CNSS listes nominatives (cnss.bj/les-listes-nominatives)"}


def _generer_csv(lignes, num_employeur, mois, annee, recap, chemin):
    import csv
    with open(chemin, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["N° employeur", num_employeur, "Mois", mois, "Année", annee,
                    "Nb travailleurs", recap["nb_travailleurs"],
                    "Masse salariale", recap["masse_salariale"]])
        w.writerow(COLONNES)
        for i, l in enumerate(lignes, start=1):
            w.writerow([i, l.num_assurance, l.nom, l.prenoms,
                        _fmt_date(l.date_embauche), _fmt_date(l.date_debauchage),
                        l.jours_ou_heures, round(l.salaire_brut),
                        round(l.rappel_salaire), l.periode_rappel,
                        round(l.indemnite_retraite), round(l.indemnite_licenciement),
                        round(l.autres_primes)])


if __name__ == "__main__":
    lignes = [
        LigneSalarie("BJ-001234", "DOSSOU", "Awa", date(2020, 3, 1),
                     jours_ou_heures=26, salaire_brut=184990),
        LigneSalarie("BJ-005678", "AGBODJAN", "Koffi", date(2018, 1, 15),
                     jours_ou_heures=26, salaire_brut=250000, autres_primes=15000),
        LigneSalarie("BJ-009012", "HOUNKPATIN", "Ren", date(2015, 6, 1),
                     date_debauchage=date(2026, 9, 30), jours_ou_heures=22,
                     salaire_brut=300000, indemnite_licenciement=510000),
    ]
    res = generer_liste_nominative(lignes, "EMP-4021", 9, 2026,
                                   "liste_nominative_cnss_demo.xlsx")
    print("Liste nominative CNSS générée :")
    for k, v in res.items():
        print(f"  {k:16} = {v}")
